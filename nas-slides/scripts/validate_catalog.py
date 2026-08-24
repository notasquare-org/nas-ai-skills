#!/usr/bin/env python3
"""Validate the non-negotiable NaS deck rules on an editable PPTX."""

from collections import Counter
from pathlib import Path
import sys
import zipfile

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE


ALLOWED_FONTS = {"Google Sans"}
ALLOWED_COLORS = {"053259", "FFFFFF", "575B61", "2C2C2C", "D8D8D6", "FB9F43"}
MIN_ANY_PT = 9
MIN_BODY_PT = 12
DIAGONAL_BAND_NAME = "NaS Diagonal Band"
DIAGONAL_BAND_ROTATION_DEG = -24.3
DIAGONAL_ROTATION_TOLERANCE_DEG = 0.5


def color_hex(color):
    try:
        return str(color.rgb).upper()
    except (AttributeError, ValueError):
        return None


def validate(path: Path):
    prs = Presentation(path)
    errors, warnings = [], []
    fonts, colors = Counter(), Counter()
    for si, slide in enumerate(prs.slides, 1):
        total_chars = 0
        orange_shapes = 0
        for shape in slide.shapes:
            if getattr(shape, "name", "") == DIAGONAL_BAND_NAME:
                if (
                    shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE
                    or shape.auto_shape_type != MSO_AUTO_SHAPE_TYPE.RECTANGLE
                ):
                    errors.append(f"slide {si}: diagonal band must be an ordinary rectangle")
                signed_rotation = (float(shape.rotation) + 180) % 360 - 180
                if abs(signed_rotation - DIAGONAL_BAND_ROTATION_DEG) > DIAGONAL_ROTATION_TOLERANCE_DEG:
                    errors.append(
                        f"slide {si}: diagonal band rotation is {signed_rotation:g}°; "
                        f"expected {DIAGONAL_BAND_ROTATION_DEG:g}°"
                    )
            if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                errors.append(f"slide {si}: native table is prohibited")
            if shape.shape_type == MSO_SHAPE_TYPE.CHART:
                errors.append(f"slide {si}: chart is prohibited")
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    total_chars += len(paragraph.text.strip())
                    for run in paragraph.runs:
                        font = run.font
                        if font.name:
                            fonts[font.name] += 1
                            if font.name not in ALLOWED_FONTS:
                                errors.append(f"slide {si}: font {font.name!r} is not Google Sans")
                        if font.size:
                            pt = font.size.pt
                            if pt < MIN_ANY_PT:
                                errors.append(f"slide {si}: text is {pt:g}pt (< {MIN_ANY_PT}pt)")
                            elif pt < MIN_BODY_PT and len(run.text.strip()) > 25:
                                warnings.append(f"slide {si}: narrative text is {pt:g}pt")
                        hx = color_hex(font.color)
                        if hx:
                            colors[hx] += 1
                            if hx not in ALLOWED_COLORS:
                                errors.append(f"slide {si}: text color #{hx} is outside palette")
            fill = getattr(shape, "fill", None)
            if fill and getattr(fill, "type", None) == MSO_FILL_TYPE.SOLID:
                hx = color_hex(fill.fore_color)
                if hx:
                    colors[hx] += 1
                    if hx not in ALLOWED_COLORS:
                        errors.append(f"slide {si}: fill color #{hx} is outside palette")
                    if hx == "FB9F43":
                        orange_shapes += 1
        if total_chars > 300:
            warnings.append(f"slide {si}: {total_chars} characters (> 300)")
        if orange_shapes > 2:
            errors.append(f"slide {si}: orange appears in {orange_shapes} shapes (> 2)")

    with zipfile.ZipFile(path) as zf:
        xml = "\n".join(
            zf.read(name).decode("utf-8", errors="ignore")
            for name in zf.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )
        if "<p:transition" in xml:
            errors.append("slide transition is prohibited")
        if "<p:timing" in xml:
            errors.append("animation timing is prohibited")
        if "<a:effectLst" in xml and "<a:effectLst/>" not in xml:
            warnings.append("effect list found; visually inspect for shadows")

    print(f"slides: {len(prs.slides)}")
    print("fonts:", dict(fonts))
    print("colors:", dict(colors))
    for item in warnings:
        print("WARNING:", item)
    for item in errors:
        print("ERROR:", item)
    return 1 if errors else 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: validate_catalog.py deck.pptx")
    raise SystemExit(validate(Path(sys.argv[1])))
